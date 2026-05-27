"""
精品採購管理系統 — 組員完整版
==============================
整合全部組員程式碼 + 下拉選單 + 各模組獨立儲存
pip install streamlit pandas openpyxl plotly numpy
"""

import streamlit as st
import pandas as pd
import numpy as np
import sqlite3, hashlib, re, io, random
import plotly.express as px
import plotly.graph_objects as go
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date, datetime, timedelta
from io import BytesIO

st.set_page_config(page_title="精品採購管理系統", page_icon="💎", layout="wide")
DB_PATH = "procurement.db"

# ============================================================
# CSS + Banner
# ============================================================
st.markdown("""
<style>
    .stApp{background-color:#f0f4f8}
    .banner{background:linear-gradient(135deg,#0a1628,#1a365d,#2b4c7e);padding:2rem;border-radius:12px;margin-bottom:1.5rem}
    .banner h1{color:#e2e8f0;text-align:center;font-size:2rem;margin:0;letter-spacing:2px}
    .banner p{color:#90a4c4;text-align:center;margin-top:.5rem}
    .card{background:white;padding:1.5rem;border-radius:12px;box-shadow:0 2px 10px rgba(0,0,0,.06);border-left:4px solid #2b4c7e;margin-bottom:1rem}
    .card h3{color:#1a365d;margin-top:0}
</style>
""", unsafe_allow_html=True)
st.markdown('<div class="banner"><h1>💎 LUXURY PROCUREMENT SYSTEM</h1><p>精品採購管理系統</p></div>', unsafe_allow_html=True)

# ============================================================
# 類別 + 範例資料
# ============================================================
big_cat_map={'RTW':'服飾 (RTW)','HB':'包款 (HB)','SLG':'小皮件 (SLG)','SH':'鞋履 (SH)','ACC':'配飾 (ACC)','JWL':'珠寶 (JWL)'}
sub_cat_map={'TP':'上衣','BT':'下著','DR':'洋裝','OW':'外套','SB':'肩背包','TT':'托特包','CB':'斜背包','CL':'手拿包','BP':'後背包','LW':'長夾','SW':'短夾','CH':'卡夾','PC':'護照套','SN':'運動鞋','LF':'樂福鞋','HL':'高跟鞋','BO':'靴子','SD':'涼鞋','SC':'絲巾','BL':'皮帶','HT':'帽子','NL':'項鍊','BC':'手鍊','RG':'戒指','ER':'耳環'}
STORES=["信義A8","忠孝SOGO","南西店","板橋店","台中新光"]
SW_MAP={"信義A8":1.5,"忠孝SOGO":1.3,"南西店":1.0,"板橋店":0.7,"台中新光":0.8}

def gen_sales():
    random.seed(42); recs=[]
    for m in range(1,13):
        d=date(2025,m,random.randint(1,28))
        for s in STORES:
            for big,subs in [('HB',['SB','TT','CB','CL']),('SH',['SN','LF','BO','SD']),('ACC',['SC','BL','HT']),('RTW',['TP','DR','OW','BT']),('SLG',['LW','SW','CH']),('JWL',['NL','BC','RG','ER'])]:
                for sub in subs:
                    q=max(1,int(10*SW_MAP[s]+random.randint(-3,5))); p=random.choice([8000,12000,18000,25000,38000,55000,78000])
                    recs.append({'門市':s,'大類別':big,'小類別':sub,'日期':str(d),'貨號':f"{big}-{sub}-{random.randint(100,999)}",'品名':sub_cat_map.get(sub,sub),
                        '顏色':random.choice(["黑色","白色","棕色","藍色","紅色"]),'尺寸':random.choice(['S','M','L','0 (ONE SIZE)']),
                        '類型':random.choice(['經典','時尚']),'實售價格':p,'數量':q,'銷售人員':f'員工{random.randint(1,20)}',
                        '會員編號':f'VIP{random.randint(1000,9999)}' if random.random()>0.3 else '',
                        '會員名稱':f'會員{random.randint(1,200)}' if random.random()>0.3 else None,
                        '會員生日':f'{random.randint(1965,2005)}-{random.randint(1,12):02d}-{random.randint(1,28):02d}' if random.random()>0.3 else None})
    df=pd.DataFrame(recs); df['銷售總額']=df['實售價格']*df['數量']; df['大類別名稱']=df['大類別'].map(big_cat_map).fillna(df['大類別'])
    df['小類別名稱']=df['小類別'].map(sub_cat_map).fillna(df['小類別']); df['日期_dt']=pd.to_datetime(df['日期'],errors='coerce')
    df['月份']=df['日期_dt'].dt.strftime('%m月'); df['類型']=df['類型'].fillna('未分類')
    if '會員生日' in df.columns:
        df['會員生日_dt']=pd.to_datetime(df['會員生日'],errors='coerce'); df['年齡']=2025-df['會員生日_dt'].dt.year
        df['年齡層']=pd.cut(df['年齡'],bins=[0,20,30,40,50,60,120],labels=['20歲以下','21-30歲','31-40歲','41-50歲','51-60歲','61歲以上'])
    return df

# ============================================================
# 資料庫
# ============================================================
def init_db():
    conn=sqlite3.connect(DB_PATH)
    conn.execute("CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, password TEXT NOT NULL, brand TEXT NOT NULL, role TEXT NOT NULL, created TEXT NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS module_data (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, module TEXT NOT NULL, date TEXT NOT NULL, summary TEXT, data_json TEXT)")
    conn.commit(); conn.close()
init_db()

def hash_pw(pw): return hashlib.sha256(pw.encode()).hexdigest()
def register(u,p,b,r):
    conn=sqlite3.connect(DB_PATH)
    try: conn.execute("INSERT INTO users (username,password,brand,role,created) VALUES (?,?,?,?,?)",(u,hash_pw(p),b,r,str(date.today()))); conn.commit(); conn.close(); return True,"註冊成功！"
    except sqlite3.IntegrityError: conn.close(); return False,"帳號已被使用"
def do_login(u,p):
    conn=sqlite3.connect(DB_PATH); cur=conn.execute("SELECT id,username,brand,role FROM users WHERE username=? AND password=?",(u,hash_pw(p))); row=cur.fetchone(); conn.close()
    return {"id":row[0],"username":row[1],"brand":row[2],"role":row[3]} if row else None
def save_module(uid,module,summary,df):
    conn=sqlite3.connect(DB_PATH); conn.execute("INSERT INTO module_data (user_id,module,date,summary,data_json) VALUES (?,?,?,?,?)",(uid,module,str(date.today()),summary,df.to_json() if df is not None and len(df)>0 else None)); conn.commit(); conn.close()
def get_module_history(uid,module):
    conn=sqlite3.connect(DB_PATH); df=pd.read_sql_query("SELECT id,date,summary,data_json FROM module_data WHERE user_id=? AND module=? ORDER BY id DESC",conn,params=(uid,module)); conn.close(); return df

def show_history(uid,module,label):
    hist=get_module_history(uid,module)
    if len(hist)==0: return
    st.divider(); st.subheader(f"📜 {label}歷史紀錄")
    for _,r in hist.iterrows():
        with st.expander(f"📅 {r['date']}　{r['summary'] or ''}"):
            if r['data_json']:
                try: st.dataframe(pd.read_json(io.StringIO(r['data_json'])),use_container_width=True,hide_index=True)
                except: st.write("資料格式無法顯示")

def df_to_xlsx(df,sheet_title="Sheet1",title_text=None):
    wb=Workbook();ws=wb.active;ws.title=sheet_title;sr=1
    hf=Font(bold=True,color="FFFFFF",name="Arial",size=11);hfill=PatternFill("solid",fgColor="1a365d")
    bdr=Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
    if title_text:
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(len(df.columns),1))
        ws.cell(row=1,column=1,value=title_text).font=Font(bold=True,size=14,name="Arial")
        ws.cell(row=2,column=1,value=f"產出日期：{date.today()}"); sr=5
    for ci,cn in enumerate(df.columns,1):
        c=ws.cell(row=sr,column=ci,value=cn);c.font=hf;c.fill=hfill;c.alignment=Alignment(horizontal="center");c.border=bdr
    for ri,row in enumerate(df.itertuples(index=False),sr+1):
        for ci,val in enumerate(row,1):
            c=ws.cell(row=ri,column=ci,value=val);c.border=bdr
            if isinstance(val,(int,float)):c.number_format="#,##0"
    for ci in range(1,len(df.columns)+1):
        ml=max(len(str(ws.cell(row=sr,column=ci).value or "")),*[len(str(ws.cell(row=r,column=ci).value or "")) for r in range(sr+1,sr+len(df)+1)])
        ws.column_dimensions[get_column_letter(ci)].width=min(ml+4,30)
    out=io.BytesIO();wb.save(out);out.seek(0);return out.getvalue()

# ============================================================
# 組員邏輯：採購數量建議（模組2）
# ============================================================
def run_procurement_model(df_sales, new_list_bytes, total_budget):
    col_style='原廠貨號';col_price='實售價格';col_qty='數量';col_season='季節';col_category='大類別';col_color='顏色編號'
    df_sales[col_qty]=pd.to_numeric(df_sales[col_qty],errors='coerce').fillna(0)
    df_sales[col_price]=pd.to_numeric(df_sales[col_price],errors='coerce').fillna(0)
    df_new=pd.read_excel(BytesIO(new_list_bytes),header=1); df_new=df_new.iloc[1:].copy()
    total_revenue=(df_sales[col_price]*df_sales[col_qty]).sum(); total_sales_qty=df_sales[col_qty].sum()
    if total_sales_qty==0: raise ValueError("歷史銷售數量為 0")
    avg_retail=total_revenue/total_sales_qty; avg_wholesale=avg_retail*0.40
    total_target=int(total_budget/avg_wholesale); co_target=int(total_target*0.60); se_target=total_target-co_target
    weight_map={'SS26':1.0,'AW26':0.8,'SS25':0.6,'AW25':0.4}
    df_sales['Weight']=df_sales[col_season].apply(lambda s:weight_map.get(str(s).strip().upper(),0.1))
    df_sales['Weighted_Qty']=df_sales[col_qty]*df_sales['Weight']
    co_styles=df_new[df_new['PLC'].str.strip().str.title()=='Carryover']['Style Number'].dropna().unique()
    df_co=df_sales[df_sales[col_style].isin(co_styles)]
    sw_sum=df_co.groupby(col_style)['Weighted_Qty'].sum(); tw=sw_sum.sum()
    co_alloc=((sw_sum/tw)*co_target).round().to_dict() if tw>0 else {}
    df_se=df_sales[~df_sales[col_style].isin(co_styles)]
    cat_sum=df_se.groupby(col_category)[col_qty].sum(); cat_prop=cat_sum/cat_sum.sum()
    cat_alloc=(cat_prop*se_target).round().to_dict()
    color_map={}
    for cat in cat_prop.index:
        top5=df_se[df_se[col_category]==cat].groupby(col_color)[col_qty].sum().nlargest(5).index
        color_map[cat]={str(c).strip():5-i for i,c in enumerate(top5)}
    df_new['Score']=0; df_new['Calculated_Order_Qty']=0
    for idx,row in df_new.iterrows():
        if str(row.get('PLC')).strip().title()=='Seasonal':
            df_new.at[idx,'Score']=color_map.get(str(row.get('Category')).strip(),{}).get(str(row.get('Color Code')).strip(),0)
    for cat,tq in cat_alloc.items():
        mask=(df_new['PLC'].str.strip().str.title()=='Seasonal')&(df_new['Category'].str.strip()==cat)
        ts=df_new.loc[mask,'Score'].sum()
        if ts>0: df_new.loc[mask,'Calculated_Order_Qty']=(df_new.loc[mask,'Score']/ts*tq).round()
        elif mask.sum()>0: df_new.loc[mask,'Calculated_Order_Qty']=round(tq/mask.sum())
    for idx,row in df_new.iterrows():
        if str(row.get('PLC')).strip().title()=='Carryover':
            df_new.at[idx,'Calculated_Order_Qty']=co_alloc.get(str(row.get('Style Number')).strip(),0)
    qty_dict=dict(zip(df_new['Style Number'].astype(str).str.strip(),df_new['Calculated_Order_Qty']))
    wb=load_workbook(BytesIO(new_list_bytes));ws=wb.active
    sc=tc=None
    for cell in ws[2]:
        if str(cell.value).strip()=='Style Number': sc=cell.column
        elif str(cell.value).strip()=='Total Order Quantity': tc=cell.column
    if not sc or not tc: raise ValueError("找不到欄位標題")
    for r in range(4,ws.max_row+1):
        v=ws.cell(row=r,column=sc).value
        if v: ws.cell(row=r,column=tc,value=int(qty_dict.get(str(v).strip(),0)))
    buf=BytesIO();wb.save(buf);wb.close();buf.seek(0)
    return buf.getvalue(),{'avg_wholesale':avg_wholesale,'total_target':total_target,'co_target':co_target,'se_target':se_target},df_new

# ============================================================
# 組員邏輯：Assortment 生成（模組3）
# ============================================================
def run_assortment(hq_bytes,assort_bytes,exchange_rate):
    from openpyxl.drawing.image import Image as XlImage
    log=[]; df=pd.read_excel(BytesIO(hq_bytes),header=1); df=df.iloc[1:].copy()
    df['Total Order Quantity']=pd.to_numeric(df['Total Order Quantity'],errors='coerce').fillna(0)
    df_f=df[df['Total Order Quantity']!=0].copy()
    if df_f.empty: raise ValueError("沒有採購數量大於 0 的商品")
    for c in ['Season','Category','Subcategory','Style Number']:
        if c in df_f.columns: df_f[c]=df_f[c].fillna('').astype(str)
    def gen_item(row):
        p=row.get('Style Number','').split('-')[0]; sc=p[-3:] if len(p)>=3 else p
        return f"{row.get('Season','')}{row.get('Category','')}{row.get('Subcategory','')}{sc}"
    df_f['Item Number']=df_f.apply(gen_item,axis=1)
    df_f['Wholesale Price(NTD)']=(pd.to_numeric(df_f.get('Wholesale Price(€)',0),errors='coerce')*exchange_rate).round()
    df_f['Retail Price(NTD)']=(pd.to_numeric(df_f.get('Suggested Retail Price(€)',0),errors='coerce')*exchange_rate).round()
    recs=df_f.to_dict('records')
    wb_in=load_workbook(BytesIO(hq_bytes));ws_in=wb_in.active
    sc_in=None
    for cell in ws_in[2]:
        if cell.value=="Style Number": sc_in=cell.column; break
    img_dict={}
    if sc_in and ws_in._images:
        for img in ws_in._images:
            ri=img.anchor._from.row+1; sid=ws_in.cell(row=ri,column=sc_in).value
            if sid: img_dict[str(sid).strip()]=img
    wb_out=load_workbook(BytesIO(assort_bytes));ws_out=wb_out.active
    existing={};scan=4
    while True:
        v=ws_out.cell(row=scan,column=2).value
        if v is None or str(v).strip()=="": break
        existing[str(v).strip()]=scan; scan+=1
    cur=scan;uc=0;ac=0
    for rd in recs:
        sid=str(rd.get('Style Number','')).strip()
        if sid in existing:
            tr=existing[sid]; eq=ws_out.cell(row=tr,column=16).value
            try: eq=float(eq) if eq else 0
            except: eq=0
            ws_out.cell(row=tr,column=16,value=eq+float(rd.get('Total Order Quantity',0)))
            es=ws_out.cell(row=tr,column=12).value; es=str(es).strip() if es else ""
            ns=str(rd.get('Season','')).strip()
            if ns and ns not in es: ws_out.cell(row=tr,column=12,value=f"{es}, {ns}" if es else ns)
            uc+=1; log.append(f"✅ 更新：{sid}")
        else:
            for ci,k in [(2,'Style Number'),(3,'Item Number'),(4,'Style Name'),(5,'Color Code'),(6,'Color Name'),(7,'Gender'),(8,'Collection'),(9,'Category'),(10,'Subcategory'),(11,'Materials'),(12,'Season'),(13,'PLC'),(14,'Wholesale Price(NTD)'),(15,'Retail Price(NTD)'),(16,'Total Order Quantity')]:
                ws_out.cell(row=cur,column=ci,value=rd.get(k,''))
            if sid in img_dict:
                try:
                    oi=img_dict[sid]; ni=XlImage(oi.ref)
                except: ni=XlImage(io.BytesIO(oi._data()))
                ratio=min(150/ni.width,150/ni.height); ni.width=int(ni.width*ratio); ni.height=int(ni.height*ratio)
                ws_out.add_image(ni,ws_out.cell(row=cur,column=1).coordinate)
            ws_out.row_dimensions[cur].height=115; cur+=1; ac+=1; log.append(f"➕ 新增：{sid}")
    ws_out.column_dimensions['A'].width=23
    buf=BytesIO();wb_out.save(buf);wb_in.close();wb_out.close();buf.seek(0)
    return buf.getvalue(),log,{'existing':len(existing),'updated':uc,'appended':ac}

# ============================================================
# 組員邏輯：配貨與調撥（模組4）
# ============================================================
def run_allocation(assort_bytes,past_bytes,current_bytes):
    store_map={'台北101':'Store1','台北微風廣場':'Store1','SOGO復興館':'Store1','板橋大遠百':'Store1','新光三越信義A4':'Store1','新竹巨城':'Store1','台中新光三越':'Store2','台中大遠百':'Store2','高雄漢神巨蛋':'Store3','高雄漢神本館':'Store3','台南新光三越西門':'Store3'}
    df_a=pd.read_excel(BytesIO(assort_bytes),header=None); df_p=pd.read_excel(BytesIO(past_bytes)); df_c=pd.read_excel(BytesIO(current_bytes))
    rn={'條碼編號':'item number','門市':'store','數量':'quantity'}
    df_p=df_p.rename(columns=rn); df_c=df_c.rename(columns=rn)
    df_p['store']=df_p['store'].map(store_map); df_c['store']=df_c['store'].map(store_map)
    if 'Order' in str(df_a.iloc[1,15]):
        df_a.insert(15,'A_S1',"");df_a.insert(16,'A_S2',"");df_a.insert(17,'A_S3',"")
        df_a=df_a.astype(object); df_a.iloc[1,15:18]=['Arrival']*3; df_a.iloc[2,15:18]=['Store1','Store2','Store3']
    else: df_a=df_a.astype(object)
    pg=df_p.groupby(['item number','store'])['quantity'].sum().unstack(fill_value=0); pg['Total']=pg.sum(axis=1)
    w1=(pg['Store1']/pg['Total']).to_dict() if 'Store1' in pg else {}
    w2=(pg['Store2']/pg['Total']).to_dict() if 'Store2' in pg else {}
    cg=df_c.groupby(['item number','store'])['quantity'].sum().unstack(fill_value=0)
    logs=[]; SAFE=2; base=datetime(2026,4,1); cnt=0
    for i in range(3,len(df_a)):
        ino=df_a.iloc[i,2]
        if pd.isna(ino) or str(ino).strip()=="": continue
        cnt+=1
        df_a.iloc[i,15]=(base+timedelta(days=int(np.random.randint(0,2)))).strftime("%Y-%m-%d")
        df_a.iloc[i,16]=(base+timedelta(days=int(np.random.randint(2,5)))).strftime("%Y-%m-%d")
        df_a.iloc[i,17]=(base+timedelta(days=int(np.random.randint(3,6)))).strftime("%Y-%m-%d")
        ot=pd.to_numeric(df_a.iloc[i,18],errors='coerce'); ot=0 if pd.isna(ot) else ot
        o1=int(round(ot*w1.get(ino,0.33)));o2=int(round(ot*w2.get(ino,0.33)));o3=int(ot-o1-o2)
        df_a.iloc[i,19:22]=[o1,o2,o3]
        s1=cg.get('Store1',pd.Series(dtype=float)).get(ino,0);s2=cg.get('Store2',pd.Series(dtype=float)).get(ino,0);s3=cg.get('Store3',pd.Series(dtype=float)).get(ino,0)
        st_=s1+s2+s3; df_a.iloc[i,22:26]=[st_,s1,s2,s3]
        inv={'Store1':o1-s1,'Store2':o2-s2,'Store3':o3-s3}
        defs={k:v for k,v in inv.items() if v<SAFE}; surs=dict(sorted({k:v for k,v in inv.items() if v>=5}.items(),key=lambda x:x[1],reverse=True))
        for ds,di in defs.items():
            need=SAFE-di
            for ss,si in list(surs.items()):
                if need<=0: break
                avail=si-SAFE
                if avail>0:
                    q=min(need,avail); logs.append({'Item Number':ino,'Style Name':df_a.iloc[i,3],'From Store':ss,'To Store':ds,'Transfer Qty':q})
                    surs[ss]-=q;inv[ss]-=q;inv[ds]+=q;need-=q
        df_a.iloc[i,30:34]=[sum(inv.values()),inv['Store1'],inv['Store2'],inv['Store3']]
        df_a.iloc[i,26:30]=[round(st_/ot,4) if ot>0 else 0,round(s1/o1,4) if o1>0 else 0,round(s2/o2,4) if o2>0 else 0,round(s3/o3,4) if o3>0 else 0]
    df_a=df_a.fillna("")
    sb=BytesIO();df_a.to_excel(sb,header=None,index=False);sb.seek(0)
    tb=None
    if logs: tb=BytesIO();pd.DataFrame(logs).to_excel(tb,index=False);tb.seek(0)
    return sb.getvalue(),(tb.getvalue() if tb else None),logs,{'items':cnt,'transfers':len(logs)}

# ============================================================
# 組員邏輯：瑕疵品（模組6a）
# ============================================================
A_kw=["包裝損傷","外盒破損","保護膜刮痕","輕微壓痕"]
B_kw=["刮傷","磨損","掉色","脫線","五金刮傷","車線歪斜","拉鍊不順","展示痕跡","小污漬","輕微掉色"]
C_kw=["破裂","斷裂","發霉","腐蝕","破損","嚴重染色","缺件","變形"]

def authenticate_item(si,sn,inu,snm,cc,cnm):
    errs=[]
    if not str(si).strip(): errs.append("Style Image empty")
    if not re.match(r"^[A-Za-z0-9\-_]{6,20}$",str(sn).strip()): errs.append("Style Number invalid")
    if not re.match(r"^[A-Za-z0-9\-_]{6,20}$",str(inu).strip()): errs.append("Item Number invalid")
    if not str(snm).strip(): errs.append("Style Name empty")
    if not re.match(r"^[A-Za-z0-9]{2,6}$",str(cc).strip()): errs.append("Color Code invalid")
    if not str(cnm).strip(): errs.append("Color Name empty")
    return (True,"OK") if not errs else (False,"; ".join(errs))

def classify_defect(si,sn,inu,snm,cc,cnm,desc,rep_cost,ws_p,ret_p,sup):
    try:
        ok,reason=authenticate_item(si,sn,inu,snm,cc,cnm)
        if not ok: return ("No","Counterfeit","Destroy",f"Auth failed: {reason}")
        desc=str(desc)
        for kw in C_kw:
            if kw in desc:
                if str(sup).strip()=="Yes": return ("Yes","Grade C","Return to Supplier",f"「{kw}」— 嚴重，供應商責任")
                return ("Yes","Grade C","Destroy",f"「{kw}」— 嚴重瑕疵")
        for kw in B_kw:
            if kw in desc:
                if ret_p==0: return ("Yes","Grade B","Pending","零售價為0")
                margin=(ret_p-ws_p-rep_cost)/ret_p
                if margin>=0.20: return ("Yes","Grade B","Repair and Sell",f"「{kw}」— 利潤率 {margin:.1%} ≥ 20%")
                return ("Yes","Grade B","Pending",f"「{kw}」— 利潤率 {margin:.1%} < 20%")
        for kw in A_kw:
            if kw in desc: return ("Yes","Grade A","Sell Normally",f"「{kw}」— 輕微包裝")
        return ("Yes","Unknown","Manual Review","未匹配關鍵字")
    except Exception as e: return ("Error","System Error","Check Data",str(e))

def run_defect(file_bytes):
    df=pd.read_excel(BytesIO(file_bytes),header=1)
    req=["Style Image","Style Number","Item Number","Style Name","Color Code","Color Name","Defect Description","Repair Cost (NTD)","Wholesale Price(NTD)","Retail Price(NTD)","Supplier Fault"]
    miss=[c for c in req if c not in df.columns]
    if miss: raise ValueError(f"缺少欄位：{miss}")
    res=df.apply(lambda r:classify_defect(r["Style Image"],r["Style Number"],r["Item Number"],r["Style Name"],r["Color Code"],r["Color Name"],r["Defect Description"],r["Repair Cost (NTD)"],r["Wholesale Price(NTD)"],r["Retail Price(NTD)"],r["Supplier Fault"]),axis=1)
    rdf=pd.DataFrame(res.tolist(),columns=["Authentic","Defect Grade","Recommended Action","Reason"])
    wb=load_workbook(BytesIO(file_bytes));ws=wb.active
    hm={str(c.value).strip():c.column for c in ws[2] if c.value}; nc=max(hm.values())+1
    for cn in ["Authentic","Defect Grade","Recommended Action","Reason"]:
        if cn not in hm: ws.cell(row=2,column=nc,value=cn);hm[cn]=nc;nc+=1
    for i,rr in rdf.iterrows():
        er=i+3
        for cn in ["Authentic","Defect Grade","Recommended Action","Reason"]: ws.cell(row=er,column=hm[cn],value=rr[cn])
    buf=BytesIO();wb.save(buf);buf.seek(0)
    return buf.getvalue(),rdf

# ============================================================
# 組員邏輯：退換貨（模組6b）
# ============================================================
def evaluate_return(si,sn,inu,snm,cc,cnm,cond,limited,price,pu,rq):
    try:
        ok,reason=authenticate_item(si,sn,inu,snm,cc,cnm)
        if not ok: return ("Rejected",f"Auth failed: {reason}")
        days=(rq-pu).days
        if days>30: return ("Rejected",f"超過退貨期限 {days} 天")
        if str(limited).strip()=="Yes": return ("Rejected","限定商品不可退")
        if cond=="Heavily Used": return ("Rejected","明顯使用痕跡")
        elif cond=="Lightly Used": return ("Exchange or Store Credit","輕微使用痕跡")
        if price>=300000: return ("Manual Review","高價需人工確認")
        return ("Return Approved","符合退貨政策")
    except Exception as e: return ("Error",str(e))

def run_return(file_bytes):
    df=pd.read_excel(BytesIO(file_bytes),header=1)
    req=["Style Image","Style Number","Item Number","Style Name","Color Code","Color Name","Product Condition","Limited Edition","Retail Price(NTD)","Purchase Date","Request Date"]
    miss=[c for c in req if c not in df.columns]
    if miss: raise ValueError(f"缺少欄位：{miss}")
    df["Purchase Date"]=pd.to_datetime(df["Purchase Date"]); df["Request Date"]=pd.to_datetime(df["Request Date"])
    res=df.apply(lambda r:evaluate_return(r["Style Image"],r["Style Number"],r["Item Number"],r["Style Name"],r["Color Code"],r["Color Name"],r["Product Condition"],r["Limited Edition"],r["Retail Price(NTD)"],r["Purchase Date"],r["Request Date"]),axis=1)
    rdf=pd.DataFrame(res.tolist(),columns=["Return Result","Reason"])
    wb=load_workbook(BytesIO(file_bytes));ws=wb.active
    hm={str(c.value).strip():c.column for c in ws[2] if c.value}; nc=max(hm.values())+1
    for cn in ["Return Result","Reason"]:
        if cn not in hm: ws.cell(row=2,column=nc,value=cn);hm[cn]=nc;nc+=1
    for i,rr in rdf.iterrows():
        er=i+3
        for cn in ["Return Result","Reason"]: ws.cell(row=er,column=hm[cn],value=rr[cn])
    buf=BytesIO();wb.save(buf);buf.seek(0)
    return buf.getvalue(),rdf

# ============================================================
# 滯銷品邏輯（模組5）
# ============================================================
def rule_engine(row,cd):
    if row.get("is_carryover",False): return "維持正價銷售"
    try: ld=datetime.strptime(str(row.get("launch_date","2026-04-01")),"%Y-%m-%d")
    except: ld=datetime(2026,4,1)
    months=(cd.year-ld.year)*12+(cd.month-ld.month)
    if months>12: return "銷毀" if row.get("brand_tier")=="Ultra_Luxury" else "送 Outlet"
    elif months>6: return "員工特賣"
    else: return "跨店調撥"

# ============================================================
# 登入 / 註冊
# ============================================================
if "user" not in st.session_state: st.session_state["user"]=None
if st.session_state["user"] is None:
    t1,t2=st.tabs(["🔑 登入","📝 註冊"])
    with t1:
        st.subheader("登入"); lu=st.text_input("帳號",key="lu"); lp=st.text_input("密碼",type="password",key="lp")
        if st.button("登入",type="primary",key="bl"):
            if lu and lp:
                r=do_login(lu,lp)
                if r: st.session_state["user"]=r;st.rerun()
                else: st.error("帳號或密碼錯誤")
    with t2:
        st.subheader("註冊"); ru=st.text_input("帳號",key="ru");rp=st.text_input("密碼",type="password",key="rp")
        rp2=st.text_input("確認密碼",type="password",key="rp2");rb=st.text_input("品牌",key="rb")
        rr=st.selectbox("角色",["採購買手 Buyer","採購主管 Manager","營運經理 Operations"],key="rr")
        if st.button("註冊",type="primary",key="br"):
            if not ru or not rp or not rb: st.warning("請填所有欄位")
            elif rp!=rp2: st.error("密碼不一致")
            elif len(rp)<4: st.error("密碼至少4字")
            else:
                ok,msg=register(ru,rp,rb,rr);st.success(msg) if ok else st.error(msg)
    st.stop()

user=st.session_state["user"]
if user is None: st.session_state.clear();st.rerun()

# ============================================================
# 導航
# ============================================================
st.sidebar.markdown(f"### 👤 {user['username']}")
st.sidebar.markdown(f"**品牌：** {user['brand']}　**角色：** {user['role']}")
st.sidebar.divider()
nav=st.sidebar.selectbox("選擇功能模組",["🏠 首頁","📊 銷售分析","📈 採購數量建議","📋 Assortment 生成","🏬 銷售概況與門市配貨","🏷️ 滯銷品處置","🔧 瑕疵品與退換貨"])
if st.sidebar.button("🚪 登出"): st.session_state.clear();st.rerun()

# ============================================================
# 首頁
# ============================================================
if nav=="🏠 首頁":
    st.markdown(f'<div class="card"><h3>歡迎回來，{user["username"]}！</h3><p>品牌：{user["brand"]}　角色：{user["role"]}</p></div>',unsafe_allow_html=True)
    st.info("從左側選單選擇功能模組開始操作。")

# ============================================================
# 1. 銷售分析
# ============================================================
elif nav=="📊 銷售分析":
    st.title("📊 銷售分析 Dashboard")
    m1=st.radio("資料來源",["上傳 Excel","使用範例資料"],key="m1r",horizontal=True)
    if m1=="上傳 Excel":
        f1=st.file_uploader("銷售 Excel（SALES 工作表）",type=["xlsx"],key="m1u")
        if f1:
            @st.cache_data
            def load_s(file):
                df=pd.read_excel(file,sheet_name='SALES');df['銷售總額']=df['實售價格']*df['數量']
                df['大類別名稱']=df['大類別'].map(big_cat_map).fillna(df['大類別']);df['小類別名稱']=df['小類別'].map(sub_cat_map).fillna(df['小類別'])
                if '日期' in df.columns: df['日期_dt']=pd.to_datetime(df['日期'],errors='coerce');df['月份']=df['日期_dt'].dt.strftime('%m月')
                if '類型' in df.columns: df['類型']=df['類型'].fillna('未分類')
                if '會員生日' in df.columns:
                    df['會員生日_dt']=pd.to_datetime(df['會員生日'],errors='coerce');df['年齡']=2025-df['會員生日_dt'].dt.year
                    df['年齡層']=pd.cut(df['年齡'],bins=[0,20,30,40,50,60,120],labels=['20歲以下','21-30歲','31-40歲','41-50歲','51-60歲','61歲以上'])
                return df
            df=load_s(f1)
        else: st.info("等待上傳..."); st.stop()
    else: df=gen_sales()

    stores=st.multiselect("篩選門市",options=df['門市'].unique(),default=df['門市'].unique())
    fdf=df[df['門市'].isin(stores)]
    if '類型' in df.columns:
        types=st.multiselect("篩選類型",options=df['類型'].unique(),default=df['類型'].unique()); fdf=fdf[fdf['類型'].isin(types)]

    c1,c2,c3,c4=st.columns(4)
    c1.metric("總銷售額",f"${fdf['銷售總額'].sum():,.0f}");c2.metric("售出",f"{fdf['數量'].sum():,.0f} 件")
    c3.metric("客單價",f"${fdf['實售價格'].mean():,.0f}");c4.metric("會員",f"{fdf[fdf['會員名稱'].notna()]['會員名稱'].nunique()} 人")
    st.divider()
    cA,cB=st.columns(2)
    with cA: st.plotly_chart(px.pie(fdf.groupby('大類別名稱')['銷售總額'].sum().reset_index(),values='銷售總額',names='大類別名稱',title='大類別佔比'),use_container_width=True)
    with cB: st.plotly_chart(px.bar(fdf.groupby('門市')['銷售總額'].sum().reset_index().sort_values('銷售總額',ascending=True),x='銷售總額',y='門市',orientation='h',title='門市排行',text_auto='.2s'),use_container_width=True)
    st.divider()
    if '銷售人員' in fdf.columns:
        st.subheader("金牌銷售員"); rep=fdf.groupby(['門市','銷售人員']).agg(銷售總額=('銷售總額','sum'),售出件數=('數量','sum')).reset_index().sort_values('銷售總額',ascending=False)
        cS1,cS2=st.columns(2)
        with cS1: t10=rep.head(10).sort_values('銷售總額',ascending=True);t10['顯示']=t10['門市']+" - "+t10['銷售人員'];st.plotly_chart(px.bar(t10,x='銷售總額',y='顯示',orientation='h',title='Top 10',text_auto='.2s'),use_container_width=True)
        with cS2: st.dataframe(rep.style.format({'銷售總額':'${:,.0f}','售出件數':'{:,.0f}'}),use_container_width=True,height=400)
        st.divider()
    if '月份' in fdf.columns:
        st.subheader("月趨勢")
        cT1,cT2=st.columns(2)
        with cT1: st.plotly_chart(px.line(fdf.groupby('月份')['銷售總額'].sum().reset_index().sort_values('月份'),x='月份',y='銷售總額',title='月銷售額',markers=True),use_container_width=True)
        with cT2:
            if '類型' in fdf.columns: st.plotly_chart(px.bar(fdf.groupby(['月份','類型'])['銷售總額'].sum().reset_index().sort_values('月份'),x='月份',y='銷售總額',color='類型',title='經典 vs 時尚',barmode='stack'),use_container_width=True)
        st.divider()
    st.subheader("小類別排行"); st.plotly_chart(px.bar(fdf.groupby(['大類別名稱','小類別名稱'])['銷售總額'].sum().reset_index().sort_values('銷售總額',ascending=False),x='小類別名稱',y='銷售總額',color='大類別名稱',text_auto='.2s'),use_container_width=True); st.divider()
    st.subheader("熱銷 TOP 30"); gcols=['大類別名稱','小類別名稱','貨號','品名']+(['類型'] if '類型' in fdf.columns else [])
    st.dataframe(fdf.groupby(gcols)['數量'].sum().reset_index().sort_values('數量',ascending=False).head(30),height=400,use_container_width=True,hide_index=True); st.divider()
    st.subheader("門市偏好"); opts=["全部門市"]+list(df['門市'].unique()); tg=st.selectbox("門市",opts); sdf=df if tg=="全部門市" else df[df['門市']==tg]
    p1,p2,p3=st.columns(3)
    with p1: st.write("**最愛品項**");st.dataframe(sdf.groupby('小類別名稱')['數量'].sum().sort_values(ascending=False).head(5))
    with p2: st.write("**最愛顏色**");st.dataframe(sdf.groupby('顏色')['數量'].sum().sort_values(ascending=False).head(5))
    with p3: st.write("**最愛尺寸**");st.dataframe(sdf[sdf['尺寸']!='0 (ONE SIZE)'].groupby('尺寸')['數量'].sum().sort_values(ascending=False).head(5))
    st.divider()
    st.subheader("VIP"); vip=fdf[fdf['會員編號'].notna()&(fdf['會員編號']!='')]
    if len(vip)>0:
        if '年齡層' in fdf.columns:
            vA,vB=st.columns([1,2])
            with vA: st.plotly_chart(px.pie(fdf[fdf['年齡層'].notna()].groupby('年齡層',observed=False)['銷售總額'].sum().reset_index(),values='銷售總額',names='年齡層',hole=0.4,title='年齡層'),use_container_width=True)
            with vB: vs=vip.groupby(['會員名稱','會員編號']).agg(總消費額=('銷售總額','sum'),購買件數=('數量','sum'),客單價=('實售價格','mean')).reset_index().sort_values('總消費額',ascending=False).head(20);st.dataframe(vs.style.format({'總消費額':'${:,.0f}','購買件數':'{:,.0f}','客單價':'${:,.0f}'}),use_container_width=True)
        else: st.dataframe(vip.groupby(['會員名稱','會員編號']).agg(總消費額=('銷售總額','sum')).reset_index().sort_values('總消費額',ascending=False).head(20),use_container_width=True,hide_index=True)
    st.divider()
    st.subheader("需求預測"); growth=st.slider("成長率 (%)",-20,50,10,1);mult=1+growth/100
    fc=fdf.groupby(['大類別名稱','小類別名稱']).agg(本年銷量=('數量','sum')).reset_index()
    fc['預測採購量']=(fc['本年銷量']*mult).astype(int);fc['增長']=fc['預測採購量']-fc['本年銷量'];fc=fc.sort_values('預測採購量',ascending=False)
    fig=go.Figure();fig.add_trace(go.Bar(x=fc['小類別名稱'],y=fc['本年銷量'],name='本年',marker_color='lightgray'))
    fig.add_trace(go.Bar(x=fc['小類別名稱'],y=fc['預測採購量'],name='預測',marker_color='#2b4c7e'));fig.update_layout(barmode='group',title="本年 vs 預測")
    st.plotly_chart(fig,use_container_width=True);st.dataframe(fc,use_container_width=True,hide_index=True)
    if st.button("💾 儲存銷售分析",type="primary"):
        st.session_state["p_sales"]=df; save_module(user["id"],"銷售分析",f"總額 ${fdf['銷售總額'].sum():,.0f}",fc); st.success("已儲存！")
    show_history(user["id"],"銷售分析","銷售分析")

# ============================================================
# 2. 採購數量建議
# ============================================================
elif nav=="📈 採購數量建議":
    st.title("📈 採購數量建議")
    st.markdown("上傳歷史銷售紀錄與新一季商品清單，系統自動計算建議採購量。")
    c1,c2=st.columns(2)
    with c1: sf=st.file_uploader("歷史銷售紀錄 Excel",type=["xlsx"],key="p2s")
    with c2: nf=st.file_uploader("新一季商品清單 Excel",type=["xlsx"],key="p2n")
    budget=st.number_input("採購總預算 (NTD)",min_value=0,value=1000000,step=10000,format="%d")
    if st.button("開始計算",type="primary",disabled=(not sf or not nf)):
        with st.spinner("分析中..."):
            try:
                ds=pd.read_excel(sf); nb=nf.read()
                ob,sm,df_r=run_procurement_model(ds,nb,budget)
                st.success("完成！")
                m1,m2,m3,m4=st.columns(4)
                m1.metric("進貨單價",f"NTD ${sm['avg_wholesale']:,.0f}");m2.metric("總目標",f"{sm['total_target']:,} 件")
                m3.metric("經典款 60%",f"{sm['co_target']:,} 件");m4.metric("時尚款 40%",f"{sm['se_target']:,} 件")
                st.divider()
                dc=[c for c in ['Style Number','PLC','Category','Color Code','Score','Calculated_Order_Qty'] if c in df_r.columns]
                st.dataframe(df_r[dc].sort_values('Calculated_Order_Qty',ascending=False),use_container_width=True,height=400)
                st.download_button("📥 下載採購建議 Excel",ob,"Auto_Forecast.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                if st.button("💾 儲存",key="sv2"):
                    save_module(user["id"],"採購建議",f"目標 {sm['total_target']:,} 件",df_r[dc]);st.success("已儲存！")
            except Exception as e: st.error(f"失敗：{e}")
    show_history(user["id"],"採購建議","採購建議")

# ============================================================
# 3. Assortment 生成
# ============================================================
elif nav=="📋 Assortment 生成":
    st.title("📋 Assortment 生成")
    st.markdown("上傳總部商品清單與 Assortment 範本，系統自動新增商品並更新既有款。")
    c1,c2=st.columns(2)
    with c1: hf=st.file_uploader("總部商品清單",type=["xlsx"],key="p3h")
    with c2: af=st.file_uploader("Assortment 範本",type=["xlsx"],key="p3a")
    fx=st.number_input("歐元兌台幣",min_value=1.0,value=35.5,step=0.1,format="%.2f")
    if st.button("開始生成",type="primary",disabled=(not hf or not af)):
        with st.spinner("處理中（含圖片）..."):
            try:
                ob,log,sm=run_assortment(hf.read(),af.read(),fx)
                st.success("完成！")
                m1,m2,m3=st.columns(3);m1.metric("原有",f"{sm['existing']} 款");m2.metric("更新",f"{sm['updated']} 款");m3.metric("新增",f"{sm['appended']} 款")
                with st.expander("處理日誌"):
                    for l in log: st.write(l)
                st.download_button("📥 下載 Assortment",ob,"Assortment_updated.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                if st.button("💾 儲存",key="sv3"):
                    save_module(user["id"],"Assortment",f"更新{sm['updated']}+新增{sm['appended']}",pd.DataFrame(log,columns=["日誌"]));st.success("已儲存！")
            except Exception as e: st.error(f"失敗：{e}")
    show_history(user["id"],"Assortment","Assortment")

# ============================================================
# 4. 配貨與調撥
# ============================================================
elif nav=="🏬 配貨與調撥":
    st.title("🏬 配貨與門市調撥")
    st.markdown("上傳三個檔案，系統自動計算配貨量並產出調撥建議。")
    c1,c2,c3=st.columns(3)
    with c1: f4a=st.file_uploader("Assortment 清單",type=["xlsx"],key="p4a")
    with c2: f4p=st.file_uploader("去年銷售紀錄",type=["xlsx"],key="p4p")
    with c3: f4c=st.file_uploader("今年銷售紀錄",type=["xlsx"],key="p4c")
    if st.button("開始配貨",type="primary",disabled=(not f4a or not f4p or not f4c)):
        with st.spinner("運算中..."):
            try:
                sb,tb,logs,sm=run_allocation(f4a.read(),f4p.read(),f4c.read())
                st.success("完成！")
                m1,m2=st.columns(2);m1.metric("品項數",f"{sm['items']} 件");m2.metric("調撥",f"{sm['transfers']} 筆")
                if logs:
                    st.subheader("調撥建議"); st.dataframe(pd.DataFrame(logs),use_container_width=True)
                st.divider()
                d1,d2=st.columns(2)
                with d1: st.download_button("📥 銷售庫存報表",sb,"output_sales_inventory.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                with d2:
                    if tb: st.download_button("📥 調撥紀錄",tb,"store_log_transfer.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                if logs and st.button("💾 儲存",key="sv4"):
                    save_module(user["id"],"配貨調撥",f"{sm['items']}品項 {sm['transfers']}筆調撥",pd.DataFrame(logs));st.success("已儲存！")
            except Exception as e: st.error(f"失敗：{e}")
    show_history(user["id"],"配貨調撥","配貨調撥")

# ============================================================
# 5. 滯銷品
# ============================================================
elif nav=="🏷️ 滯銷品處置":
    st.title("🏷️ 滯銷品處置")
    st.markdown("上傳 Assortment（含庫存）分析滯銷品處置方式。")
    f5=st.file_uploader("Assortment / 庫存 Excel",type=["xlsx"],key="p5u")
    if f5:
        try:
            df5=pd.read_excel(f5,header=1); df5=df5.iloc[1:].copy()
            cd=datetime.now(); results=[]
            for _,row in df5.iterrows():
                sn=str(row.get('Style Number','')); plc=str(row.get('PLC','Seasonal')).strip()
                ic=plc.lower()=="carryover"; ret_p=float(row.get('Retail Price(NTD)',0) or 0); ws_p=float(row.get('Wholesale Price(NTD)',0) or 0)
                action=rule_engine({"is_carryover":ic,"launch_date":"2026-04-01","brand_tier":"Luxury","current_location":"正價專門店"},cd)
                est=ret_p if action=="維持正價銷售" else int(ret_p*0.5) if action=="送 Outlet" else int(ret_p*0.3) if action=="員工特賣" else 0
                loss=0 if action=="維持正價銷售" else ret_p-est if action!="銷毀" else ws_p
                results.append({"Style Number":sn,"Style Name":row.get('Style Name',''),"PLC":plc,"處置建議":action,"預估回收":est,"預估損失":loss})
            df_r=pd.DataFrame(results)
            slow=df_r[df_r["處置建議"]!="維持正價銷售"]
            if len(slow)>0:
                st.metric("需處置",f"{len(slow)} 項"); st.dataframe(slow,use_container_width=True,hide_index=True)
                sm_=slow.groupby("處置建議").agg(品項=("Style Number","count"),預估回收=("預估回收","sum"),預估損失=("預估損失","sum")).reset_index()
                st.dataframe(sm_,use_container_width=True,hide_index=True)
                st.download_button("📥 滯銷品清單",df_to_xlsx(slow,"滯銷","滯銷品處置清單"),"滯銷品.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                if st.button("💾 儲存",key="sv5"): save_module(user["id"],"滯銷品",f"需處置 {len(slow)} 項",slow);st.success("已儲存！")
            else: st.success("🎉 無滯銷品！")
        except Exception as e: st.error(f"失敗：{e}")
    else: st.info("等待上傳...")
    show_history(user["id"],"滯銷品","滯銷品")

# ============================================================
# 6. 瑕疵品 & 退換貨
# ============================================================
elif nav=="🔧 瑕疵品與退換貨":
    st.title("🔧 瑕疵品處理 & 退換貨評估")
    tab1,tab2=st.tabs(["🔍 瑕疵品分類","↩️ 退換貨評估"])

    with tab1:
        st.subheader("瑕疵品自動分類")
        with st.expander("查看判定規則"):
            ca,cb,cc=st.columns(3)
            with ca: st.markdown("**Grade A**"); [st.markdown(f"- {k}") for k in A_kw]
            with cb: st.markdown("**Grade B**"); [st.markdown(f"- {k}") for k in B_kw]
            with cc: st.markdown("**Grade C**"); [st.markdown(f"- {k}") for k in C_kw]
        f6d=st.file_uploader("瑕疵品 Excel",type=["xlsx"],key="p6d")
        if st.button("開始分類",type="primary",disabled=not f6d,key="b6d"):
            with st.spinner("分析中..."):
                try:
                    ob,rdf=run_defect(f6d.read()); st.success(f"完成，{len(rdf)} 筆")
                    # 用 set 統計不重複等級
                    unique_grades=set(rdf["Defect Grade"].tolist()); st.caption(f"偵測到等級：{unique_grades}")
                    gc=rdf["Defect Grade"].value_counts()
                    cols=st.columns(len(gc))
                    for i,(g,c) in enumerate(gc.items()): cols[i].metric(g,f"{c} 件")
                    if (rdf["Authentic"]=="No").sum()>0: st.error(f"⚠️ {(rdf['Authentic']=='No').sum()} 件疑似仿冒！")
                    st.divider()
                    grade_colors={"Grade A":"background-color:#d4edda","Grade B":"background-color:#fff3cd","Grade C":"background-color:#f8d7da","Counterfeit":"background-color:#6c1212;color:white","Unknown":"background-color:#e2e3e5"}
                    st.dataframe(rdf.style.apply(lambda r:[grade_colors.get(r["Defect Grade"],"")]*len(r),axis=1),use_container_width=True,height=350)
                    st.download_button("📥 瑕疵品結果",ob,"瑕疵品分類結果.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    if st.button("💾 儲存瑕疵品",key="sv6d"): save_module(user["id"],"瑕疵品",f"{len(rdf)} 筆",rdf);st.success("已儲存！")
                except Exception as e: st.error(f"失敗：{e}")
        show_history(user["id"],"瑕疵品","瑕疵品")

    with tab2:
        st.subheader("退換貨資格評估")
        with st.expander("查看判定規則"):
            st.markdown("| 條件 | 結果 |\n|---|---|\n| 驗證失敗 | Rejected |\n| >30天 | Rejected |\n| 限定版 | Rejected |\n| 嚴重使用 | Rejected |\n| 輕微使用 | Exchange |\n| ≥30萬 | Manual Review |\n| 其他 | Approved |")
        f6r=st.file_uploader("退換貨 Excel",type=["xlsx"],key="p6r")
        if st.button("開始評估",type="primary",disabled=not f6r,key="b6r"):
            with st.spinner("評估中..."):
                try:
                    ob,rdf=run_return(f6r.read()); st.success(f"完成，{len(rdf)} 筆")
                    rc=rdf["Return Result"].value_counts()
                    cols=st.columns(len(rc))
                    for i,(r,c) in enumerate(rc.items()): cols[i].metric(r,f"{c} 筆")
                    st.divider()
                    ret_colors={"Return Approved":"background-color:#d4edda","Exchange or Store Credit":"background-color:#cce5ff","Manual Review":"background-color:#fff3cd","Rejected":"background-color:#f8d7da"}
                    st.dataframe(rdf.style.apply(lambda r:[ret_colors.get(r["Return Result"],"")]*len(r),axis=1),use_container_width=True,height=350)
                    st.download_button("📥 退換貨結果",ob,"退換貨結果.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    if st.button("💾 儲存退換貨",key="sv6r"): save_module(user["id"],"退換貨",f"{len(rdf)} 筆",rdf);st.success("已儲存！")
                except Exception as e: st.error(f"失敗：{e}")
        show_history(user["id"],"退換貨","退換貨")
